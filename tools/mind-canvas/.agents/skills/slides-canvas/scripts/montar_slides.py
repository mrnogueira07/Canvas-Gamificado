#!/usr/bin/env python3
"""
Mind Canvas · Slides Canvas
Monta o .pptx de uma videoaula a partir do JSON gerado pelo agente,
usando o template oficial de cada linha (Konectar.me, Konecta Kids,
Educação Corporativa). Opcionalmente exporta PDF com o LibreOffice.

Uso:
    python montar_slides.py saida/AULA.json
    python montar_slides.py saida/AULA.json --pdf

Requer: pip install python-pptx
"""
import argparse
import copy
import json
import math
import os
import shutil
import subprocess
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

AQUI = os.path.dirname(os.path.abspath(__file__))
TEMPLATES = os.path.join(AQUI, "..", "resources", "templates")

PRETO = "000000"
BRANCO = "FFFFFF"

# ---------------------------------------------------------------------------
# Configuração de cada linha. Medidas em polegadas, fontes em pontos.
# "area_util" = retângulo onde o conteúdo pode ficar (fora da área de segurança).
# ---------------------------------------------------------------------------
LINHAS = {
    "Konectar.me": {
        "template": "konectar-me.pptx",
        "base": {"capa": 1, "conteudo": 2, "destaque": 2, "final": 11},
        "fonte_titulo": "Montserrat SemiBold", "tam_titulo": 100, "cor_titulo": "2796E1",
        "fonte_texto": "Montserrat Medium", "tam_texto": 80, "tam_texto_longo": 60,
        "tam_fonte_img": 40, "fonte_formula": "Montserrat", "tam_formula": 100,
        "cores_separacao": ["F79646", "7030A0", "9BBB59"], "contorno": 10,
        "area_util": (3.0, 3.0, 36.0, 25.0),          # x, y, largura, altura (3/4 esquerdos)
        "capa": [
            {"rotulo": "Tema da aula:", "campo": "tema", "x": 20.9, "y": 8.7},
            {"rotulo": "Disciplina:", "campo": "disciplina", "x": 20.9, "y": 16.8},
        ],
        "capa_fonte": "Montserrat SemiBold", "capa_tam": 100, "capa_cor_rotulo": "257CE1",
        "capa_cor_valor": PRETO, "capa_larg": 25.5, "capa_passo": 1.7,
        "rotulos_objetivo": ("OBJETIVO", "HABILIDADE BNCC"),
    },
    "Konecta Kids": {
        "template": "konecta-kids.pptx",
        "base": {"capa": 1, "conteudo": 2, "destaque": 3, "final": 4},
        "fonte_titulo": "Lexend Medium", "tam_titulo": 100, "cor_titulo": "571075",
        "fonte_texto": "Lexend", "tam_texto": 80, "tam_texto_longo": 60,
        "tam_fonte_img": 40, "fonte_formula": "Lexend Medium", "tam_formula": 100,
        "cores_separacao": ["993DBC", "32DAC3", "FFC500", "FE647C"], "contorno": 10,
        "area_util": (5.2, 3.6, 34.8, 21.5),
        "capa": [
            {"rotulo": None, "campo": "professor", "x": 6.6, "y": 12.5},
            {"rotulo": None, "campo": "tema", "x": 6.6, "y": 17.3},
        ],
        "capa_fonte": "Lexend Medium", "capa_tam": 100, "capa_cor_rotulo": "571075",
        "capa_cor_valor": PRETO, "capa_larg": 38.0, "capa_passo": 0,
        "rotulos_objetivo": ("Objetivo:", "Habilidade da BNCC:"),
    },
    "Educação Corporativa": {
        "template": "educacao-corporativa.pptx",
        "base": {"capa": 1, "conteudo": 2, "destaque": 2, "final": 3},
        "fonte_titulo": "Montserrat Medium", "titulo_negrito": True, "tam_titulo": 28, "cor_titulo": "03153D",
        "fonte_texto": "Montserrat Medium", "tam_texto": 18, "tam_texto_longo": 14,
        "tam_fonte_img": 10, "fonte_formula": "Montserrat", "tam_formula": 28,
        "cores_separacao": ["6FAADB", "03153D"], "contorno": 2.5,
        "area_util": (0.6, 0.6, 8.9, 6.3),             # 3/4 esquerdos, entre a linha do topo e a faixa da base
        "capa": [
            {"rotulo": None, "campo": "titulo", "x": 0.6, "y": 2.7, "tam": 28, "negrito": True, "cor": "6FAADB"},
            {"rotulo": None, "campo": "subtitulo", "x": 0.6, "y": 3.25, "tam": 18, "negrito": False, "cor": BRANCO},
        ],
        "capa_fonte": "Montserrat Medium", "capa_tam": 28, "capa_cor_rotulo": "6FAADB",
        "capa_cor_valor": BRANCO, "capa_larg": 6.5, "capa_passo": 0,
        "rotulos_objetivo": ("Objetivo", "Competência desenvolvida"),
    },
}
APELIDOS = {"konectar.me": "Konectar.me", "konectarme": "Konectar.me", "konectar": "Konectar.me",
            "kids": "Konecta Kids", "konecta kids": "Konecta Kids",
            "corporativa": "Educação Corporativa", "educação corporativa": "Educação Corporativa",
            "educacao corporativa": "Educação Corporativa"}

avisos = []


def aviso(msg):
    avisos.append(msg)


# ---------------------------------------------------------------------------
# Clonar um slide-base do template (fundo + imagens), sem as caixas de texto.
# ---------------------------------------------------------------------------
def clonar_base(prs, src):
    novo = prs.slides.add_slide(src.slide_layout)
    for shp in list(novo.shapes):
        shp._element.getparent().remove(shp._element)
    mapa = {}
    for rid, rel in src.part.rels.items():
        if rel.is_external:
            continue
        if "image" in rel.reltype or "media" in rel.reltype:
            mapa[rid] = novo.part.relate_to(rel._target, rel.reltype)

    def remapear(el):
        for node in el.iter():
            for attr in list(node.attrib):
                if attr.endswith("}embed") or attr.endswith("}link") or attr.endswith("}id"):
                    if node.attrib[attr] in mapa:
                        node.attrib[attr] = mapa[node.attrib[attr]]
        return el

    src_csld = src._element.cSld
    novo_csld = novo._element.cSld
    if src_csld.bg is not None:
        bg = remapear(copy.deepcopy(src_csld.bg))
        if novo_csld.bg is not None:
            novo_csld.remove(novo_csld.bg)
        novo_csld.insert(0, bg)
    # copia só elementos decorativos: imagens e grupos sem texto
    for shp in src.shapes:
        tem_texto = shp.has_text_frame and shp.text_frame.text.strip()
        if shp.shape_type in (13, 6) and not tem_texto:          # PICTURE, GROUP
            novo.shapes._spTree.append(remapear(copy.deepcopy(shp._element)))
    return novo


def remover_slides_originais(prs, n):
    lista = prs.slides._sldIdLst
    for sld in list(lista)[:n]:
        prs.part.drop_rel(sld.rId)
        lista.remove(sld)


# ---------------------------------------------------------------------------
# Texto
# ---------------------------------------------------------------------------
def altura_estimada(texto, larg_in, tam_pt):
    por_linha = max(1, int(larg_in * 72 / (tam_pt * 0.56)))
    linhas = sum(max(1, math.ceil(len(p) / por_linha)) for p in texto.split("\n"))
    return linhas * tam_pt * 1.25 / 72, linhas


def caixa(slide, x, y, w, h, paragrafos, fonte, tam, cor=PRETO, negrito=False,
          alinhar=PP_ALIGN.LEFT, ancora=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancora
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(0.05))
    if isinstance(paragrafos, str):
        paragrafos = [paragrafos]
    for i, t in enumerate(paragrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinhar
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = fonte
        f.size = Pt(tam)
        f.bold = negrito
        f.color.rgb = RGBColor.from_string(cor)
    return tb


def retangulo(slide, x, y, w, h, contorno=None, larg_contorno=2, preenchimento=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if preenchimento:
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor.from_string(preenchimento)
    else:
        s.fill.background()
    if contorno:
        s.line.color.rgb = RGBColor.from_string(contorno)
        s.line.width = Pt(larg_contorno)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


# ---------------------------------------------------------------------------
# Tipos de slide
# ---------------------------------------------------------------------------
def slide_capa(slide, cfg, dados):
    for item in cfg["capa"]:
        valor = dados.get(item["campo"]) or f"[{item['campo'].upper()}]"
        tam = item.get("tam", cfg["capa_tam"])
        if item["rotulo"]:
            caixa(slide, item["x"], item["y"], cfg["capa_larg"], tam / 72 * 1.3, item["rotulo"],
                  cfg["capa_fonte"], tam, cfg["capa_cor_rotulo"], negrito=True)
            y_valor = item["y"] + cfg["capa_passo"]
        else:
            y_valor = item["y"]
        caixa(slide, item["x"], y_valor, cfg["capa_larg"], tam / 72 * 1.3, valor,
              cfg["capa_fonte"], tam, item.get("cor", cfg["capa_cor_valor"]),
              negrito=item.get("negrito", False))


def titulo(slide, cfg, texto, cor=None):
    x, y, w, _ = cfg["area_util"]
    tam = cfg["tam_titulo"]
    h, _ = altura_estimada(texto, w, tam)
    caixa(slide, x, y, w, h, texto, cfg["fonte_titulo"], tam, cor or cfg["cor_titulo"],
          negrito=cfg.get("titulo_negrito", False))
    return y + h + tam / 72 * 0.6


def slide_objetivo(slide, cfg, s):
    x, y, w, hmax = cfg["area_util"]
    r1, r2 = cfg["rotulos_objetivo"]
    blocos = [(r1, s.get("objetivo") or " ".join(s.get("texto", [])[:1])),
              (r2, s.get("habilidade") or " ".join(s.get("texto", [])[1:]))]
    tam_t, tam = cfg["tam_titulo"], cfg["tam_texto"]
    total = sum(altura_estimada(t, w, tam)[0] for _, t in blocos)
    if total + 2 * tam_t / 72 * 1.6 > hmax:
        tam = cfg["tam_texto_longo"]
        aviso(f"Slide {s['numero']}: texto longo, fonte reduzida para {tam} pt (ponto a decidir no manual).")
    for rot, txt in blocos:
        hr = tam_t / 72 * 1.3
        caixa(slide, x, y, w, hr, rot, cfg["fonte_titulo"], tam_t, cfg["cor_titulo"],
              negrito=cfg.get("titulo_negrito", False))
        y += hr + 0.15
        h, _ = altura_estimada(txt, w, tam)
        caixa(slide, x, y, w, h, txt, cfg["fonte_texto"], tam)
        y += h + tam_t / 72 * 0.9


def bloco_imagem(slide, cfg, img, x, y, w, h):
    arq = (img or {}).get("arquivo")
    if arq and os.path.exists(arq):
        slide.shapes.add_picture(arq, Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        r = retangulo(slide, x, y, w, h, contorno="7F7F7F", larg_contorno=2, preenchimento="EDEDED")
        tf = r.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = "[BUSCAR IMAGEM] " + (img or {}).get("descricao", "")
        run.font.size = Pt(cfg["tam_fonte_img"])
        run.font.name = cfg["fonte_texto"]
        run.font.color.rgb = RGBColor.from_string("404040")
        aviso(f"Imagem pendente: {(img or {}).get('descricao', '(sem descrição)')}")
    fonte = (img or {}).get("fonte") or "Fonte: (autor) / (banco de imagens)."
    tam = cfg["tam_fonte_img"]
    caixa(slide, x, y + h + 0.1, w, tam / 72 * 1.4, fonte, cfg["fonte_texto"], tam,
          alinhar=PP_ALIGN.CENTER)


def slide_conteudo(slide, cfg, s, cor_titulo=None, cor_texto=PRETO):
    x0, y0, w, hmax = cfg["area_util"]
    y = titulo(slide, cfg, s.get("titulo", ""), cor_titulo) if s.get("titulo") else y0
    limite = y0 + hmax
    tam = cfg["tam_texto"]
    itens = [t for t in s.get("texto", []) if t]
    layout = s.get("layout", "so_texto")
    max_itens = 3 if cfg is LINHAS["Konecta Kids"] else 5
    if len(itens) > max_itens:
        aviso(f"Slide {s['numero']}: {len(itens)} itens; o limite da linha é {max_itens}. Divida o slide.")

    if layout == "imagem_esquerda_texto_direita" or (s.get("imagem") and layout == "so_texto"):
        wi = w * 0.45
        hi = min(limite - y - 2 * cfg["tam_fonte_img"] / 72, wi * 0.62)
        bloco_imagem(slide, cfg, s.get("imagem"), x0, y, wi, hi)
        xt, wt = x0 + wi + w * 0.05, w * 0.5
        txt = "\n".join(itens)
        h, _ = altura_estimada(txt, wt, tam)
        caixa(slide, xt, y, wt, max(h, hi), itens, cfg["fonte_texto"], tam, cor_texto)
        return

    if layout == "formula_com_espaco":
        f = s.get("formula", "")
        hf = cfg["tam_formula"] / 72 * 1.4
        caixa(slide, x0, y, w, hf, f, cfg["fonte_formula"], cfg["tam_formula"], negrito=True)
        y += hf + 0.3
        retangulo(slide, x0, y, w * 0.9, max(1.0, limite - y), contorno=cfg["cores_separacao"][0],
                  larg_contorno=cfg["contorno"])
        return

    if layout in ("topicos_com_barras", "caixas_contorno"):
        cores = cfg["cores_separacao"]
        gap = tam / 72 * 0.8
        for i, t in enumerate(itens):
            cor = cores[i % len(cores)]
            h, _ = altura_estimada(t, w * 0.9, tam)
            if layout == "topicos_com_barras":
                retangulo(slide, x0, y, max(0.08, w * 0.012), h, preenchimento=cor)
                caixa(slide, x0 + w * 0.03, y, w * 0.95, h, t, cfg["fonte_texto"], tam, cor_texto)
            else:
                retangulo(slide, x0, y, w, h + 0.2, contorno=cor, larg_contorno=cfg["contorno"])
                caixa(slide, x0 + 0.1, y + 0.1, w - 0.2, h, t, cfg["fonte_texto"], tam, cor_texto)
            y += h + gap
        if y > limite:
            aviso(f"Slide {s['numero']}: o texto passa da área útil. Reduza ou divida.")
        return

    # so_texto: tópicos com marcador quando houver mais de um item
    paras = [("• " + t) if len(itens) > 1 else t for t in itens]
    h, linhas = altura_estimada("\n".join(paras), w, tam)
    caixa(slide, x0, y, w, max(h, 0.5), paras, cfg["fonte_texto"], tam, cor_texto)
    if y + h > limite:
        aviso(f"Slide {s['numero']}: o texto passa da área útil. Reduza ou divida.")


def notas(slide, texto):
    if texto:
        slide.notes_slide.notes_text_frame.text = texto


# ---------------------------------------------------------------------------
def montar(caminho_json, pdf=False):
    with open(caminho_json, encoding="utf-8") as fh:
        dados = json.load(fh)
    nome_linha = APELIDOS.get(str(dados.get("linha", "")).strip().lower(), dados.get("linha"))
    if nome_linha not in LINHAS:
        sys.exit(f"Linha desconhecida: {dados.get('linha')}. Use Konectar.me, Konecta Kids ou Educação Corporativa.")
    cfg = LINHAS[nome_linha]
    prs = Presentation(os.path.join(TEMPLATES, cfg["template"]))
    originais = list(prs.slides)
    n_orig = len(originais)

    for s in dados["slides"]:
        tipo = s.get("tipo", "conteudo")
        if tipo == "capa":
            base = "capa"
        elif tipo == "encerramento":
            base = "final"
        elif s.get("fundo") == "destaque":
            base = "destaque"
        else:
            base = "conteudo"
        slide = clonar_base(prs, originais[cfg["base"][base] - 1])
        if tipo == "capa":
            slide_capa(slide, cfg, {**dados.get("capa", {}), **s})
        elif tipo == "encerramento":
            pass
        elif tipo == "objetivo_bncc":
            slide_objetivo(slide, cfg, s)
        elif base == "destaque" and nome_linha == "Konecta Kids":
            slide_conteudo(slide, cfg, s, cor_titulo=BRANCO, cor_texto=BRANCO)
        else:
            slide_conteudo(slide, cfg, s)
        notas(slide, s.get("nota_apresentador"))
        for p in s.get("pendencias", []) or []:
            aviso(f"Slide {s['numero']}: {p}")

    remover_slides_originais(prs, n_orig)
    pasta = os.path.dirname(os.path.abspath(caminho_json))
    nome = dados.get("arquivo") or os.path.splitext(os.path.basename(caminho_json))[0]
    destino = os.path.join(pasta, nome + ".pptx")
    prs.save(destino)
    print(f"OK: {destino} ({len(dados['slides'])} slides, linha {nome_linha})")

    if pdf:
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if soffice:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", pasta, destino],
                           check=False, capture_output=True)
            print(f"PDF: {os.path.splitext(destino)[0]}.pdf")
        else:
            aviso("LibreOffice não encontrado: exporte o PDF pelo PowerPoint ou Google Slides.")

    tempo = sum(int(s.get("tempo_seg", 0) or 0) for s in dados["slides"])
    if tempo:
        print(f"Tempo estimado: {tempo // 60} min {tempo % 60:02d} s")
    if avisos:
        print("\nPENDÊNCIAS E AVISOS:")
        for a in avisos:
            print(" -", a)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Monta o .pptx da videoaula no template da linha.")
    ap.add_argument("json", help="Arquivo JSON gerado pela skill slides-canvas")
    ap.add_argument("--pdf", action="store_true", help="Exporta também em PDF (requer LibreOffice)")
    a = ap.parse_args()
    montar(a.json, a.pdf)
